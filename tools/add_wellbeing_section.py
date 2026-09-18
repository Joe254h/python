# -*- coding: utf-8 -*-
"""Insert a WHO-5 wellbeing section before Section J (Close), all in red,
with first-person notes the author can use when briefing the client."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = ('/root/.claude/uploads/920ad5d8-6e07-5909-9e17-40c3ac0fe128/'
       '3645ea1e-Trader_Questionnaire_Integrated_edited_final_final.docx')
OUT = '/home/user/python/output/Trader_Questionnaire_with_Wellbeing.docx'

F = 'Calibri'
RED = RGBColor(0xC0, 0x00, 0x00)
doc = Document(SRC)
body = doc.element.body
# lxml hands out a fresh proxy object per access, so id() is not stable —
# count the content elements instead. python-docx appends before <w:sectPr>.
def _content(b):
    return [el for el in b if el.tag != qn('w:sectPr')]
n_before = len(_content(body))

def red(r, size=11, bold=False, italic=False):
    r.font.name = F; r.font.size = Pt(size)
    r.font.bold = bold; r.font.italic = italic; r.font.color.rgb = RED
    rpr = r._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), F)
    return r

def para(text='', size=11, bold=False, italic=False, before=0, after=4, indent=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    if text:
        red(p.add_run(text), size, bold, italic)
    return p

def h1(text):
    p = para(text, size=15, bold=True, before=16, after=8)
    p.paragraph_format.keep_with_next = True
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement('w:pBdr'); bo = OxmlElement('w:bottom')
    bo.set(qn('w:val'), 'single'); bo.set(qn('w:sz'), '6')
    bo.set(qn('w:space'), '4'); bo.set(qn('w:color'), 'C00000')
    b.append(bo); pPr.append(b)
    return p

def h2(text):
    p = para(text, size=12.5, bold=True, before=12, after=5)
    p.paragraph_format.keep_with_next = True
    return p

def q(num, text, codes=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5); p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Cm(1.05)
    p.paragraph_format.first_line_indent = Cm(-1.05)
    p.paragraph_format.keep_with_next = True
    red(p.add_run(num + '  '), bold=True)
    red(p.add_run(text))
    if codes:
        c = doc.add_paragraph()
        c.paragraph_format.space_before = Pt(0); c.paragraph_format.space_after = Pt(3)
        c.paragraph_format.left_indent = Cm(1.05)
        red(c.add_run(codes), size=10.5)
    return p

def instr(text):
    return para(text, size=10, italic=True, before=2, after=4, indent=1.05)

def mynote(text):
    """First-person note the author can read straight out to the client."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Cm(0.5)
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement('w:pBdr'); e = OxmlElement('w:left')
    e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '18')
    e.set(qn('w:space'), '8'); e.set(qn('w:color'), 'C00000')
    b.append(e); pPr.append(b)
    red(p.add_run('Note on this addition:  '), size=10.5, bold=True)
    red(p.add_run(text), size=10.5)
    return p

def shade(cell, hexval):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexval)
    tcPr.append(sh)

def cell(c, text, bold=False, size=9.5, align=None):
    c.text = ''
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
    if align is not None:
        p.alignment = align
    if text:
        red(p.add_run(text), size=size, bold=bold)
    return p

def grid(header, rows, widths, size=9.5, hdrsize=9):
    t = doc.add_table(rows=len(rows) + 1, cols=len(header))
    t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.LEFT; t.autofit = False
    for r in t.rows:
        for i, w in enumerate(widths):
            r.cells[i].width = Cm(w)
    for i, lab in enumerate(header):
        cell(t.rows[0].cells[i], lab, bold=True, size=hdrsize, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(t.rows[0].cells[i], 'FBE4E4')
    trPr = t.rows[0]._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader'); el.set(qn('w:val'), 'true'); trPr.append(el)
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            cell(t.rows[ri].cells[ci], val, bold=(ci == 0 and val != ''), size=size)
    return t

# ============================================================ NEW SECTION J
h1('Section J — How you have been feeling')
para('This section is new. I have printed it in red so you can see exactly what I added and '
     'where.', size=10, italic=True, after=8)

mynote('I added this because we are asking the County to spend money on childcare, and at the '
       'moment the only benefit we can show is trading days recovered. Minding a small child at '
       'a stall all day while trying to sell is wearing, and if that strain eases when childcare '
       'is available, that is part of the return the County is buying. I chose the WHO-5 because '
       'it is five questions, takes about a minute, is free to use, and is worded positively — '
       'nothing in it sounds like a diagnosis, which matters when we are interviewing at a stall '
       'with other traders standing nearby.')

mynote('I put it here, near the end, on purpose. It comes after the choice cards and after the '
       'willingness-to-pay questions, because asking someone to dwell on how they have been '
       'feeling and then asking what they would pay would move the answer. It also sits well '
       'after the last of the childcare questions is done.')

mynote('One thing I want to be straight with you about: we are surveying each trader once. That '
       'means we can show whether traders under more strain use childcare differently — we '
       'cannot show that childcare improved anyone\'s wellbeing. If you want that second claim, '
       'we would need to measure the same traders again once a service is running. I would '
       'rather agree that now than have it questioned when the report is read.')

h2('Before you ask these questions')
instr('Read the introduction below word for word. Keep your tone ordinary — the same as for the '
      'rest of the questionnaire. If the respondent does not want to answer, move on without '
      'comment and code 9.')
para('“The last few questions are about how you have been feeling recently, not about childcare. '
     'As with everything else, you can skip any of them.”', size=11, italic=True, after=6,
     indent=1.05)

h2('J1.  WHO-5 Well-Being Index')
instr('Read: “Please say which answer comes closest to how you have felt over the last two '
      'weeks.” Read all six options aloud for the first statement, then point to the card for '
      'the rest. Tick one box per row. All five rows must be answered or the score cannot be '
      'calculated.')
who5 = [
    ('a', 'I have felt cheerful and in good spirits'),
    ('b', 'I have felt calm and relaxed'),
    ('c', 'I have felt active and vigorous'),
    ('d', 'I woke up feeling fresh and rested'),
    ('e', 'My daily life has been filled with things that interest me'),
]
scale = ['5\nAll of\nthe time', '4\nMost of\nthe time', '3\nMore than\nhalf the time',
         '2\nLess than\nhalf the time', '1\nSome of\nthe time', '0\nAt no\ntime']
t = doc.add_table(rows=len(who5) + 1, cols=8)
t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.LEFT; t.autofit = False
widths = [0.7, 5.5, 1.65, 1.65, 1.75, 1.75, 1.65, 1.6]
for r in t.rows:
    for i, w in enumerate(widths):
        r.cells[i].width = Cm(w)
for i, lab in enumerate(['', 'Over the last two weeks …'] + scale):
    cell(t.rows[0].cells[i], lab, bold=True, size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    shade(t.rows[0].cells[i], 'FBE4E4')
trPr = t.rows[0]._tr.get_or_add_trPr()
el = OxmlElement('w:tblHeader'); el.set(qn('w:val'), 'true'); trPr.append(el)
for i, (ltr, txt) in enumerate(who5, start=1):
    cell(t.rows[i].cells[0], ltr, bold=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    cell(t.rows[i].cells[1], txt, size=9.5)
    for j in range(2, 8):
        cell(t.rows[i].cells[j], '☐', size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
instr('Code 9 = declined to answer. Use the exact wording above — it is a standard instrument '
      'and changing the words breaks comparability with every other study that uses it.')

mynote('I did not reword any of the five statements, and I would ask that we keep it that way. '
       'The value of using a standard measure is that our numbers can be read against other '
       'studies; if we reword it, it becomes our own scale and that comparison is gone. For the '
       'same reason we should use an existing Kiswahili version rather than translating it '
       'ourselves, and I will source one before the pilot.')

h2('J2.  Scoring — office use, not read aloud')
grid(['Step', 'How'],
     [['J2a.  Raw score', 'Add the five answers. Range 0 to 25.'],
      ['J2b.  Percentage score', 'Multiply the raw score by 4. Range 0 to 100.'],
      ['J2c.  Flag for follow-up',
       'A percentage score of 50 or below is the usual point at which further assessment is '
       'suggested. Confirm the current threshold with the County health department before '
       'fieldwork and record which threshold we used.'],
      ['J2d.  Incomplete', 'If any of the five rows is missing, do not compute a score.']],
     widths=[4.0, 12.2])

h2('Offering support')
instr('Offer the referral card to anyone whose score is 50 or below, to anyone who becomes '
      'distressed, and to anyone who asks. Hand it over without comment and without singling '
      'the person out. Offer it the same way every time.')
q('J3.', 'Referral card offered?', '1 = Yes    2 = No — respondent declined it    '
                                   '3 = Not applicable')
q('J4.', 'Did the respondent become distressed at any point in this section?  '
         '(enumerator judgement)', '1 = Yes    2 = No')
instr('If yes: stop this section, do not probe, offer the card, and note it in your observations '
      'at the close. Do not record any detail the respondent gave about their situation.')

mynote('This is the part I want to raise with you directly. Once we ask these questions, some '
       'people will tell us they are struggling, and we cannot simply write the answer down and '
       'move on. Before we go to the field I need a referral point agreed with the County health '
       'department — somewhere real, reachable from the market, and free — printed on a card the '
       'enumerators carry. I also need to add a line to the consent script saying some questions '
       'ask how the respondent has been feeling and can be skipped, and the ethics submission '
       'has to say we are collecting this and how we store it. That submission is the longest '
       'lead item, so if we are doing this, I would like to start it now. If the timeline will '
       'not carry it, I would rather drop this section than run it without the referral in '
       'place.')

h2('J5.  Consent wording to be added')
instr('Add this sentence to the consent script, after the sentence about some questions feeling '
      'personal:')
para('“Towards the end I will ask five short questions about how you have been feeling in '
     'yourself over the last two weeks. You can skip them, and if you would like, I can give you '
     'the details of somewhere you can talk to someone.”', size=11, italic=True, after=6,
     indent=1.05)

# ---- move everything built above to sit just before "Section J — Close" ----
new_elems = _content(body)[n_before:]
target = None
for p in doc.paragraphs:
    if p.text.strip() == 'Section J — Close':
        target = p._p
        break
if target is None:
    raise SystemExit('could not find "Section J — Close"')

pb = OxmlElement('w:p')                      # page break before the new section
ppr = OxmlElement('w:pPr'); pb.append(ppr)
r = OxmlElement('w:r'); br = OxmlElement('w:br')
br.set(qn('w:type'), 'page'); r.append(br); pb.append(r)
target.addprevious(pb)
for el in new_elems:
    body.remove(el)
    target.addprevious(el)

# ---- relabel the old Close section, in red, and say why ----
for run in list(target.findall(qn('w:r'))):
    target.remove(run)
from docx.text.paragraph import Paragraph
tp = Paragraph(target, doc)
red(tp.add_run('Section K — Close'), size=15, bold=True)

note = OxmlElement('w:p')
target.addnext(note)
np_ = Paragraph(note, doc)
np_.paragraph_format.space_before = Pt(2); np_.paragraph_format.space_after = Pt(8)
red(np_.add_run('I relabelled this section from J to K, and its questions from J to K, because '
                'the wellbeing questions now sit before it. Not a word of the questions '
                'themselves changed — only the letter in front of them.'), size=10, italic=True)

# ---- renumber the Close section's own items J -> K, marking each in red ----
# Word splits "J4." across several runs, so match on the paragraph text and
# recolour every run that carries part of the label.
import re
seen = []
started = False
for par in doc.paragraphs:
    if par._p is target:
        started = True
        continue
    if not started or not par.runs:
        continue
    m = re.match(r'^(J\d+[a-z]?\.)(\s|$)', par.text)
    if not m:
        continue
    label = m.group(1)
    consumed = 0
    for r in par.runs:
        if consumed >= len(label):
            break
        if consumed == 0 and r.text.startswith('J'):
            r.text = 'K' + r.text[1:]
        red(r, size=r.font.size.pt if r.font.size else 11, bold=True)
        consumed += len(r.text)
    seen.append('K' + label[1:])
print('renumbered:', seen)

doc.save(OUT)
print('saved', OUT)
