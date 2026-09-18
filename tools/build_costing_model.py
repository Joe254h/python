# -*- coding: utf-8 -*-
"""Childcare costing model: four delivery models compared on one set of drivers."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = '/home/user/python/output/Childcare_Costing_Model.xlsx'

FONT = 'Arial'
BLUE  = Font(name=FONT, size=10, color='0000FF')            # hardcoded input
BLACK = Font(name=FONT, size=10)                            # formula
GREEN = Font(name=FONT, size=10, color='008000')            # link to another sheet
BOLD  = Font(name=FONT, size=10, bold=True)
BOLDW = Font(name=FONT, size=10, bold=True, color='FFFFFF')
TITLE = Font(name=FONT, size=14, bold=True, color='1F3864')
SUB   = Font(name=FONT, size=10, italic=True, color='595959')
NOTE  = Font(name=FONT, size=9, italic=True, color='595959')
YELLOW = PatternFill('solid', fgColor='FFFF00')
HDR    = PatternFill('solid', fgColor='1F4D78')
BAND   = PatternFill('solid', fgColor='DEEAF6')
thin = Side(style='thin', color='BFBFBF')
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
TOPLINE = Border(top=Side(style='thin', color='1F4D78'))

KES = '#,##0;(#,##0);-'
KES2 = '#,##0.00;(#,##0.00);-'
PCT = '0.0%'
NUM = '#,##0.0;(#,##0.0);-'

MODELS = ['County-run', 'Private operator', 'Partnership (PPP)', 'Trader committee']
MCOL = ['C', 'D', 'E', 'F']          # model columns used on every sheet

wb = Workbook()

def sheet(name, widths):
    ws = wb.create_sheet(name)
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.sheet_view.showGridLines = False
    return ws

def title(ws, text, sub=None):
    ws['A1'] = text; ws['A1'].font = TITLE
    if sub:
        ws['A2'] = sub; ws['A2'].font = SUB
    ws.freeze_panes = 'A5'

def hrow(ws, row, labels, start=1):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=start + i, value=lab)
        c.font = BOLDW; c.fill = HDR; c.border = BOX
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[row].height = 30

def label(ws, row, text, bold=False, indent=0):
    c = ws.cell(row=row, column=1, value=text)
    c.font = BOLD if bold else BLACK
    if indent:
        c.alignment = Alignment(indent=indent)
    return c

def inp(ws, cell, value, fmt=KES, key=False):
    c = ws[cell]; c.value = value; c.font = BLUE; c.number_format = fmt; c.border = BOX
    if key:
        c.fill = YELLOW
    return c

def fml(ws, cell, formula, fmt=KES, bold=False, link=False):
    c = ws[cell]; c.value = formula
    c.font = (BOLD if bold else (GREEN if link else BLACK))
    c.number_format = fmt; c.border = BOX
    return c

def note(ws, row, text, col=1):
    c = ws.cell(row=row, column=col, value=text); c.font = NOTE
    return c

# ================================================================== README
ws = wb.active; ws.title = 'README'
ws.column_dimensions['A'].width = 4
ws.column_dimensions['B'].width = 104
ws.sheet_view.showGridLines = False
ws['B1'] = 'Childcare Costing Model'; ws['B1'].font = TITLE
ws['B2'] = ('Uasin Gishu County markets — comparing four delivery models on cost, '
            'break-even and subsidy')
ws['B2'].font = SUB
lines = [
    ('h', 'What this does'),
    ('p', 'It takes the cost figures collected on Section L of the questionnaire and works out, '
          'for each of the four delivery models, four things: what a child-day costs, how many '
          'children are needed to break even, how large a subsidy the County would have to '
          'carry, and what inclusion costs on its own.'),
    ('h', 'How to use it'),
    ('p', '1.  Fill in the Assumptions sheet. Every blue cell is an input. Yellow cells are the '
          'ones that move the answer most — get those right first.'),
    ('p', '2.  Fill in the Costs and Revenue sheets, one column per model. Blue cells only.'),
    ('p', '3.  Read Comparison. Nothing there is typed; it is all calculated.'),
    ('p', '4.  Read Sensitivity to see whether the ranking survives being wrong about the '
          'assumptions.'),
    ('h', 'Colour code'),
    ('p', 'Blue text = a number you type in.     Black = calculated, do not overwrite.     '
          'Green = pulled from another sheet.     Yellow fill = an assumption worth arguing '
          'about.'),
    ('h', 'The figures currently in the workbook are placeholders'),
    ('p', 'They are there so the formulas compute and you can see the expected format. They are '
          'plausible orders of magnitude for a Kenyan county market, not researched figures. '
          'Every one must be replaced with a real number from the facility assessment, the '
          'county interview, a provider interview, or a supplier quotation before any result '
          'is quoted to anyone.'),
    ('h', 'The one thing not to get wrong'),
    ('p', 'County-run will look cheapest if the building it already owns and the ECDE staff '
          'already on payroll are entered as free. They are not free — they are costs the '
          'County is already carrying. The Assumptions sheet has two switches, "Charge rent to '
          'the county model" and "Charge full cost for seconded staff". Both default to Yes. '
          'Turning either off makes the comparison unfair, and the recommendation that comes '
          'out of it wrong.'),
    ('h', 'What break-even assumes'),
    ('p', 'Break-even splits costs into fixed and variable. Staff are treated as fixed, but in '
          'reality they step up each time enrolment crosses the caregiver ratio, so the figure '
          'is an approximation that is most accurate near the enrolment you entered. Read it '
          'alongside the Sensitivity sheet rather than on its own.'),
]
r = 4
for kind, text in lines:
    c = ws.cell(row=r, column=2, value=text)
    if kind == 'h':
        c.font = Font(name=FONT, size=11, bold=True, color='1F4D78')
        r += 1
    else:
        c.font = Font(name=FONT, size=10)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[r].height = 14 * (len(text) // 95 + 1)
        r += 2

# ================================================================== ASSUMPTIONS
ws = sheet('Assumptions', {'A': 46, 'B': 14, 'C': 17, 'D': 17, 'E': 17, 'F': 17, 'G': 46})
title(ws, 'Assumptions', 'Blue cells are inputs. Yellow marks the ones that move the answer most.')

ws['A4'] = 'Shared drivers'; ws['A4'].font = BOLD
hrow(ws, 5, ['Driver', 'Value', '', '', '', '', 'Where this comes from'])
shared = [
    ('Days open per year',               260,   NUM,  True,  'Trading days the centre operates. Section L6e.'),
    ('Children enrolled per day',        40,    NUM,  True,  'Expected enrolment, not places. Section L6b.'),
    ('Attendance rate',                  0.80,  PCT,  True,  'Share of enrolled children present on a given day. L6c.'),
    ('Fee charged per child per day',    60,    KES,  True,  'The price being tested. Cross-check against G1/G2.'),
    ('Fee collection rate',              0.85,  PCT,  True,  'Share of fees actually collected. G10 and G11.'),
    ('Caregiver to child ratio (1 to N)', 10,   NUM,  True,  'Enforced minimum from the county interview, section 1.'),
    ('Food cost per child per day',      35,    KES,  False, 'Porridge and lunch. L5a.'),
    ('Consumables per child per day',    8,     KES,  False, 'Soap, cleaning, first aid. L5d and L5e.'),
    ('Statutory on-costs (NSSF, SHA, leave)', 0.18, PCT, False, 'Applied to all salaries. Confirm current rates.'),
    ('Building works life (years)',      15,    NUM,  False, 'For spreading one-off works. L2.'),
    ('Equipment and training life (years)', 5,  NUM,  False, 'Furniture, play materials, initial training.'),
    ('Levy per trader per month',        0,     KES,  False, 'Only if a market-fee levy is used. G12a.'),
    ('Number of traders in the market',  1200,  NUM,  False, 'From the county market fee register.'),
    ('County subsidy per month',         0,     KES,  False, 'Any direct county contribution.'),
]
r = 6
ASSUM = {}
for name, val, fmt, key, src in shared:
    label(ws, r, name)
    inp(ws, f'B{r}', val, fmt, key=key)
    note(ws, r, src, col=7)
    ASSUM[name] = f'Assumptions!$B${r}'
    r += 1

r += 1
ws.cell(row=r, column=1, value='Imputation switches — 1 = Yes, 0 = No').font = BOLD
r += 1
label(ws, r, 'Charge rent to the county model')
inp(ws, f'B{r}', 1, NUM, key=True)
note(ws, r, 'Leave at 1. Setting 0 treats a county-owned building as free and rigs the '
            'comparison.', col=7)
SW_RENT = f'Assumptions!$B${r}'; r += 1
label(ws, r, 'Charge full cost for seconded staff')
inp(ws, f'B{r}', 1, NUM, key=True)
note(ws, r, 'Leave at 1. Seconded ECDE staff are a real cost the County already carries.', col=7)
SW_STAFF = f'Assumptions!$B${r}'; r += 2

ws.cell(row=r, column=1, value='By model').font = BOLD
r += 1
hrow(ws, r, ['Driver'] + [''] + MODELS + ['Where this comes from'])
MHDR = r
r += 1
permodel = [
    ('Rent applies (1 = yes, 0 = no)',      [1, 1, 1, 0],       NUM,
     'County pays imputed rent only if the switch above is on; the formula handles it.'),
    ('Rent per month (market rate)',        [25000, 25000, 25000, 0], KES,
     'What the space would let for. L5f.'),
    ('Caregiver pay, per month',            [22000, 18000, 20000, 14000], KES,
     'County on ECDE scale; private at market rate; committee typically lowest. L4a.'),
    ('Supervisor pay, per month',           [35000, 30000, 32000, 0],  KES, 'L4c.'),
    ('Cook pay, per month',                 [14000, 13000, 13500, 12000], KES, 'L4d.'),
    ('Cleaner pay, per month',              [12000, 11000, 11500, 10000], KES, 'L4e.'),
    ('Security pay, per month',             [15000, 14000, 14500, 0],  KES, 'L4f.'),
    ('Admin overhead, % of staff cost',     [0.12, 0.10, 0.11, 0.04],  PCT,
     'County overhead is absorbed elsewhere — impute it, do not zero it. L5j.'),
    ('Operator margin, % of total cost',    [0, 0.12, 0.06, 0],        PCT,
     'A private operator will not run at cost. L8d.'),
    ('Cost of collecting fees, % of fees',  [0.02, 0.02, 0.02, 0.06],  PCT,
     'Highest where traders collect from each other. L5k.'),
    ('Utilities per month',                 [6000, 6000, 6000, 6000],  KES, 'L5b and L5c.'),
    ('Maintenance per year',                [40000, 40000, 40000, 25000], KES, 'L5g.'),
    ('Insurance per year',                  [30000, 45000, 38000, 20000], KES,
     'Public liability. Usually required for licensing. L5h.'),
    ('Licence renewal per year',            [8000, 8000, 8000, 8000],  KES, 'L5i.'),
]
PM = {}
for name, vals, fmt, src in permodel:
    label(ws, r, name)
    for i, v in enumerate(vals):
        inp(ws, f'{MCOL[i]}{r}', v, fmt)
    note(ws, r, src, col=7)
    PM[name] = r
    r += 1
note(ws, r + 1, 'Every figure above is a placeholder. Replace each one before quoting any '
                'result.', col=1)

A = ASSUM
DAYS, CHILD, ATT = A['Days open per year'], A['Children enrolled per day'], A['Attendance rate']
FEE, COLL = A['Fee charged per child per day'], A['Fee collection rate']
RATIO = A['Caregiver to child ratio (1 to N)']
FOOD, CONS = A['Food cost per child per day'], A['Consumables per child per day']
STAT = A['Statutory on-costs (NSSF, SHA, leave)']
LIFE_W, LIFE_E = A['Building works life (years)'], A['Equipment and training life (years)']
LEVY, TRADERS = A['Levy per trader per month'], A['Number of traders in the market']
SUBSIDY = A['County subsidy per month']
P = lambda n, i: f"Assumptions!${MCOL[i]}${PM[n]}"

# ================================================================== COSTS
ws = sheet('Costs', {'A': 48, 'B': 3, 'C': 16, 'D': 16, 'E': 16, 'F': 16, 'G': 44})
title(ws, 'Costs', 'One column per model. Blue cells are inputs; black cells calculate.')
hrow(ws, 5, ['Cost line', ''] + MODELS + ['Note'])

R = {}
r = 6
def block(name):
    global r
    c = ws.cell(row=r, column=1, value=name); c.font = BOLD; c.fill = BAND
    for col in ['B'] + MCOL + ['G']:
        ws[f'{col}{r}'].fill = BAND
    r += 1

def cost_inputs(rows):
    global r
    for name, vals, src in rows:
        label(ws, r, name, indent=1)
        for i, v in enumerate(vals):
            inp(ws, f'{MCOL[i]}{r}', v)
        note(ws, r, src, col=7)
        R[name] = r
        r += 1

def total_row(name, first, last, fmt=KES, src=''):
    global r
    label(ws, r, name, bold=True)
    for col in MCOL:
        fml(ws, f'{col}{r}', f'=SUM({col}{first}:{col}{last})', fmt, bold=True)
        ws[f'{col}{r}'].border = TOPLINE
    if src: note(ws, r, src, col=7)
    R[name] = r; r += 1
    return R[name]

block('One-off capital costs — works (L2)')
start = r
cost_inputs([
    ('Building repair or conversion', [450000, 450000, 450000, 450000], 'L2a. From the facility assessment.'),
    ('Water connection or tank',      [80000, 80000, 80000, 80000],     'L2b.'),
    ('Toilets and handwashing',       [180000, 180000, 180000, 180000], 'L2c.'),
    ('Kitchen / food preparation',    [120000, 120000, 120000, 90000],  'L2d.'),
    ('Safety works',                  [90000, 90000, 90000, 70000],     'L2e.'),
    ('Electrical works and lighting', [70000, 70000, 70000, 70000],     'L2f.'),
])
W_END = r - 1
total_row('Subtotal — works', start, W_END)
SUB_W = R['Subtotal — works']

block('One-off capital costs — equipment and setup (L2)')
start = r
cost_inputs([
    ('Furniture, mats, cots, play materials', [150000, 150000, 150000, 110000], 'L2g.'),
    ('Initial caregiver training',            [120000, 120000, 120000, 60000],  'L2h.'),
    ('Registration and licensing fees',       [25000, 25000, 25000, 25000],     'L2i.'),
])
total_row('Subtotal — equipment and setup', start, r - 1)
SUB_E = R['Subtotal — equipment and setup']

block('Making the space accessible (L3) — kept separate on purpose')
start = r
cost_inputs([
    ('Ramp',                          [60000, 60000, 60000, 60000], 'L3a.'),
    ('Widening doors to 800mm',       [35000, 35000, 35000, 35000], 'L3b.'),
    ('Accessible toilet',             [95000, 95000, 95000, 95000], 'L3c.'),
    ('Levelling the route from stalls', [55000, 55000, 55000, 55000], 'L3d.'),
])
total_row('Subtotal — inclusion, one-off', start, r - 1, src='Feeds L3h and the Comparison sheet.')
SUB_I = R['Subtotal — inclusion, one-off']

label(ws, r, 'Annualised capital cost', bold=True)
for col in MCOL:
    fml(ws, f'{col}{r}',
        f'=({col}{SUB_W}+{col}{SUB_I})/{LIFE_W}+{col}{SUB_E}/{LIFE_E}', KES, bold=True)
note(ws, r, 'Works and accessibility spread over the building life; equipment and training '
            'over the shorter life.', col=7)
ANN_CAP = r; r += 2

block('Annual staff cost (L4)')
label(ws, r, 'Caregivers needed (children ÷ ratio, rounded up)', indent=1)
for col in MCOL:
    fml(ws, f'{col}{r}', f'=ROUNDUP({CHILD}/{RATIO},0)', NUM)
note(ws, r, 'Steps up each time enrolment crosses the ratio.', col=7)
N_CG = r; r += 1

label(ws, r, 'Extra caregiver time for inclusion (full-time equivalents)', indent=1)
for i, v in enumerate([0.5, 0.5, 0.5, 0.5]):
    inp(ws, f'{MCOL[i]}{r}', v, NUM)
note(ws, r, 'L3f. A lower ratio for children who need more support.', col=7)
N_INC = r; r += 1

start = r
label(ws, r, 'Caregivers', indent=1)
for i, col in enumerate(MCOL):
    fml(ws, f'{col}{r}', f'={col}{N_CG}*{P("Caregiver pay, per month",i)}*12*(1+{STAT})')
note(ws, r, 'L4a.', col=7)
CG_ROW = r; r += 1
for nm, key in [('Supervisor', 'Supervisor pay, per month'), ('Cook', 'Cook pay, per month'),
                ('Cleaner', 'Cleaner pay, per month'), ('Security', 'Security pay, per month')]:
    label(ws, r, nm, indent=1)
    for i, col in enumerate(MCOL):
        fml(ws, f'{col}{r}', f'={P(key,i)}*12*(1+{STAT})')
    r += 1
label(ws, r, 'Inclusion — extra caregiver time', indent=1)
for i, col in enumerate(MCOL):
    fml(ws, f'{col}{r}', f'={col}{N_INC}*{P("Caregiver pay, per month",i)}*12*(1+{STAT})')
note(ws, r, 'Part of the inclusion cost; also counted in the inclusion line below.', col=7)
INC_STAFF = r; r += 1
total_row('Subtotal — staff', start, r - 1)
SUB_S = R['Subtotal — staff']

label(ws, r, 'Seconded-staff adjustment', bold=True)
for i, col in enumerate(MCOL):
    if i == 0:
        fml(ws, f'{col}{r}', f'=IF({SW_STAFF}=1,0,-{col}{SUB_S})')
    else:
        fml(ws, f'{col}{r}', '=0')
note(ws, r, 'Zero while the switch is on, which is how it should stay. If the switch is turned '
            'off, county staff cost drops out entirely — that is what makes the comparison '
            'unfair.', col=7)
ADJ_S = r; r += 2

block('Other annual running costs (L5)')
start = r
label(ws, r, 'Food', indent=1)
for col in MCOL:
    fml(ws, f'{col}{r}', f'={CHILD}*{ATT}*{DAYS}*{FOOD}')
note(ws, r, 'L5a. Varies with the number of children.', col=7)
FOOD_ROW = r; r += 1
label(ws, r, 'Consumables, cleaning and first aid', indent=1)
for col in MCOL:
    fml(ws, f'{col}{r}', f'={CHILD}*{ATT}*{DAYS}*{CONS}')
CONS_ROW = r; r += 1
label(ws, r, 'Utilities', indent=1)
for i, col in enumerate(MCOL):
    fml(ws, f'{col}{r}', f'={P("Utilities per month",i)}*12')
r += 1
label(ws, r, 'Rent or space charge', indent=1)
for i, col in enumerate(MCOL):
    if i == 0:
        fml(ws, f'{col}{r}',
            f'=IF({SW_RENT}=1,{P("Rent per month (market rate)",i)}*12*{P("Rent applies (1 = yes, 0 = no)",i)},0)')
    else:
        fml(ws, f'{col}{r}',
            f'={P("Rent per month (market rate)",i)}*12*{P("Rent applies (1 = yes, 0 = no)",i)}')
note(ws, r, 'L5f. The county model pays imputed rent while the switch is on.', col=7)
r += 1
for nm, key in [('Maintenance and repairs', 'Maintenance per year'),
                ('Insurance (public liability)', 'Insurance per year'),
                ('Licence renewal', 'Licence renewal per year')]:
    label(ws, r, nm, indent=1)
    for i, col in enumerate(MCOL):
        fml(ws, f'{col}{r}', f'={P(key,i)}')
    r += 1
label(ws, r, 'Administration and management', indent=1)
for i, col in enumerate(MCOL):
    fml(ws, f'{col}{r}', f'={P("Admin overhead, % of staff cost",i)}*({SUB_S if False else f"{col}{SUB_S}"})')
r += 1
label(ws, r, 'Cost of collecting fees', indent=1)
for i, col in enumerate(MCOL):
    fml(ws, f'{col}{r}',
        f'={P("Cost of collecting fees, % of fees",i)}*{CHILD}*{ATT}*{DAYS}*{FEE}*{COLL}')
note(ws, r, 'L5k. Highest under the committee model.', col=7)
r += 1
total_row('Subtotal — other running costs', start, r - 1)
SUB_O = R['Subtotal — other running costs']

label(ws, r, 'Operator margin', bold=True)
for i, col in enumerate(MCOL):
    fml(ws, f'{col}{r}',
        f'={P("Operator margin, % of total cost",i)}*({col}{ANN_CAP}+{col}{SUB_S}+{col}{ADJ_S}+{col}{SUB_O})')
note(ws, r, 'L8d. A private operator will not run at cost.', col=7)
MARGIN = r; r += 1

label(ws, r, 'TOTAL ANNUAL COST', bold=True)
for col in MCOL:
    c = fml(ws, f'{col}{r}',
            f'={col}{ANN_CAP}+{col}{SUB_S}+{col}{ADJ_S}+{col}{SUB_O}+{col}{MARGIN}', KES, bold=True)
    c.border = TOPLINE; c.fill = BAND
TOTAL_COST = r; r += 2

label(ws, r, 'Of which, cost of inclusion', bold=True)
for col in MCOL:
    fml(ws, f'{col}{r}', f'={col}{SUB_I}/{LIFE_W}+{col}{INC_STAFF}', KES, bold=True)
note(ws, r, 'L3h annualised plus L3i. Shown separately so the County can see what inclusion '
            'costs and decide who pays for it.', col=7)
INC_TOTAL = r

# ================================================================== REVENUE
ws = sheet('Revenue', {'A': 48, 'B': 3, 'C': 16, 'D': 16, 'E': 16, 'F': 16, 'G': 44})
title(ws, 'Revenue', 'Money coming in each year, per model (Section L7).')
hrow(ws, 5, ['Source', ''] + MODELS + ['Note'])
r = 6
label(ws, r, 'Child-days per year', bold=True)
for col in MCOL:
    fml(ws, f'{col}{r}', f'={CHILD}*{ATT}*{DAYS}', NUM, bold=True)
note(ws, r, 'Children enrolled × attendance × days open. The denominator for cost per child-day.',
     col=7)
CHILD_DAYS = r; r += 2

start = r
label(ws, r, 'Fees from parents', indent=1)
for col in MCOL:
    fml(ws, f'{col}{r}', f'={col}{CHILD_DAYS}*{FEE}*{COLL}')
note(ws, r, 'L7a. Net of what is not collected — see the collection rate on Assumptions.', col=7)
r += 1
label(ws, r, 'Market levy', indent=1)
for col in MCOL:
    fml(ws, f'{col}{r}', f'={LEVY}*{TRADERS}*12')
note(ws, r, 'L7c. Zero unless a levy is used. G12a gives what traders would accept.', col=7)
r += 1
label(ws, r, 'County subsidy', indent=1)
for col in MCOL:
    fml(ws, f'{col}{r}', f'={SUBSIDY}*12')
note(ws, r, 'L7d.', col=7)
r += 1
label(ws, r, 'Donor or NGO support', indent=1)
for i in range(4):
    inp(ws, f'{MCOL[i]}{r}', 0)
note(ws, r, 'L7e.', col=7)
r += 1
label(ws, r, 'TOTAL ANNUAL REVENUE', bold=True)
for col in MCOL:
    c = fml(ws, f'{col}{r}', f'=SUM({col}{start}:{col}{r-1})', KES, bold=True)
    c.border = TOPLINE; c.fill = BAND
TOTAL_REV = r

# ================================================================== COMPARISON
ws = sheet('Comparison', {'A': 50, 'B': 3, 'C': 16, 'D': 16, 'E': 16, 'F': 16, 'G': 44})
title(ws, 'Comparison', 'Nothing on this sheet is typed. Change the inputs, not these cells.')
hrow(ws, 5, ['Result', ''] + MODELS + ['How it is worked out'])
r = 6
def crow(name, f, fmt=KES, src='', bold=True, band=False):
    global r
    label(ws, r, name, bold=bold)
    for col in MCOL:
        c = fml(ws, f'{col}{r}', f(col), fmt, bold=bold)
        if band:
            c.fill = BAND
    if src: note(ws, r, src, col=7)
    out = r; r += 1
    return out

crow('Total annual cost', lambda c: f"=Costs!{c}{TOTAL_COST}", src='From the Costs sheet.')
crow('Total annual revenue', lambda c: f"=Revenue!{c}{TOTAL_REV}", src='From the Revenue sheet.')
r += 1
C_CPD = crow('Cost per child-day', lambda c: f"=IFERROR(Costs!{c}{TOTAL_COST}/Revenue!{c}{CHILD_DAYS},0)",
             KES2, 'L9a. Total annual cost ÷ child-days. The number to compare models on.',
             band=True)
C_GAP = crow('Annual shortfall the County must cover',
             lambda c: f"=Costs!{c}{TOTAL_COST}-Revenue!{c}{TOTAL_REV}", KES,
             'L9c. Cost minus revenue. This is the budget line.', band=True)
crow('Shortfall per child-day',
     lambda c: f"=IFERROR((Costs!{c}{TOTAL_COST}-Revenue!{c}{TOTAL_REV})/Revenue!{c}{CHILD_DAYS},0)",
     KES2, 'The same gap, per child served per day.', bold=False)
r += 1

# fixed / variable split for break-even
label(ws, r, 'Variable cost per child-day', bold=False)
for col in MCOL:
    fml(ws, f'{col}{r}', f'={FOOD}+{CONS}', KES2)
note(ws, r, 'Food and consumables — the costs that rise with each extra child.', col=7)
VAR_CPD = r; r += 1
label(ws, r, 'Fixed cost per year', bold=False)
for col in MCOL:
    fml(ws, f'{col}{r}',
        f'=Costs!{col}{TOTAL_COST}-(Costs!{col}{FOOD_ROW}+Costs!{col}{CONS_ROW})', KES)
note(ws, r, 'Everything else: staff, rent, utilities, insurance, overhead, annualised capital.',
     col=7)
FIX = r; r += 1
label(ws, r, 'Net fee received per child-day', bold=False)
for col in MCOL:
    fml(ws, f'{col}{r}', f'={FEE}*{COLL}-{col}{VAR_CPD}', KES2)
note(ws, r, 'What one extra child-day contributes towards fixed costs.', col=7)
CONTRIB = r; r += 1

C_BE = crow('Children per day needed to break even',
            lambda c: (f"=IF({c}{CONTRIB}<=0,\"never at this fee\","
                       f"ROUNDUP({c}{FIX}/({c}{CONTRIB}*{DAYS}*{ATT}),0))"),
            NUM, 'L9b. If the fee does not cover food and consumables, no enrolment breaks '
                 'even — the cell says so.', band=True)
crow('Break-even fee per child-day at the enrolment entered',
     lambda c: f"=IFERROR(Costs!{c}{TOTAL_COST}/(Revenue!{c}{CHILD_DAYS}*{COLL}),0)", KES2,
     'What you would have to charge to cover costs at the enrolment on the Assumptions sheet. '
     'Compare against G1 and G2.', bold=False)
r += 1
crow('Cost of inclusion, per year', lambda c: f"=Costs!{c}{INC_TOTAL}", KES,
     'L9d. Accessibility works spread over the building life, plus the extra caregiver time.')
crow('Cost of inclusion, per child-day',
     lambda c: f"=IFERROR(Costs!{c}{INC_TOTAL}/Revenue!{c}{CHILD_DAYS},0)", KES2,
     'What it would add to the fee if inclusion were funded by all parents rather than the '
     'County. Compare against H16a.', bold=False)
r += 2
ws.cell(row=r, column=1,
        value='Read the break-even figure alongside the Sensitivity sheet. Staff costs step up '
              'when enrolment crosses the caregiver ratio, so a break-even far from the '
              'enrolment you entered is only approximate.').font = NOTE

# ================================================================== SENSITIVITY
ws = sheet('Sensitivity', {'A': 42, 'B': 13, 'C': 13, 'D': 16, 'E': 16, 'F': 16, 'G': 16,
                           'H': 3, 'I': 36})
title(ws, 'Sensitivity',
      'Does the cheapest model stay the cheapest when the assumptions are wrong?')
ws['A4'] = ('Each row changes one thing and recomputes cost per child-day. '
            'Multipliers are inputs — change them to test your own scenarios.')
ws['A4'].font = SUB

hrow(ws, 6, ['Scenario', 'Enrolment ×', 'Fixed cost ×'] + MODELS + ['', 'What it tests'])
scenarios = [
    ('As entered (base case)',            1.00, 1.00, 'The Assumptions sheet as it stands.'),
    ('Enrolment 30% lower',               0.70, 1.00, 'Fewer children than hoped — the most common way these fail.'),
    ('Enrolment 30% higher',              1.30, 1.00, 'Demand exceeds the plan.'),
    ('Salaries and overheads 20% higher', 1.00, 1.20, 'Pay pressure, or an underestimated staffing need.'),
    ('Low season',                        0.60, 1.00, 'Traders trade fewer days in some months (B21); fixed costs do not fall.'),
    ('Both: low enrolment and higher costs', 0.70, 1.20, 'The pessimistic case. If the ranking holds here, it holds.'),
]
r = 7
first_s = r
for name, em, fm, why in scenarios:
    label(ws, r, name)
    inp(ws, f'B{r}', em, '0.00')
    inp(ws, f'C{r}', fm, '0.00')
    for i, col in enumerate(['D', 'E', 'F', 'G']):
        mc = MCOL[i]
        fml(ws, f'{col}{r}',
            f'=IFERROR((Comparison!${mc}${FIX}*$C{r})/({CHILD}*$B{r}*{ATT}*{DAYS})'
            f'+Comparison!${mc}${VAR_CPD},0)', KES2)
    note(ws, r, why, col=9)
    r += 1
last_s = r - 1

r += 1
ws.cell(row=r, column=1, value='Cheapest model in each scenario').font = BOLD
r += 1
for i, (name, _, _, _) in enumerate(scenarios):
    row_i = first_s + i
    label(ws, r, name, indent=1)
    fml(ws, f'B{r}',
        f'=INDEX({{"{MODELS[0]}","{MODELS[1]}","{MODELS[2]}","{MODELS[3]}"}},'
        f'MATCH(MIN(D{row_i}:G{row_i}),D{row_i}:G{row_i},0))', 'General')
    ws[f'B{r}'].alignment = Alignment(horizontal='left')
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    r += 1
r += 1
ws.cell(row=r, column=1,
        value='If the same name appears on every row, the recommendation is robust. If it '
              'changes, say so plainly and tell the County which figure decides it.').font = NOTE
r += 1
ws.cell(row=r, column=1,
        value='Simplification: this sheet scales fixed costs by a single multiplier rather than '
              'rebuilding the staffing roster. It is a test of robustness, not a second model — '
              'for a specific scenario, change the Assumptions sheet itself and read '
              'Comparison.').font = NOTE

wb.save(OUT)
print('saved', OUT)
