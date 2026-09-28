# -*- coding: utf-8 -*-
"""Update the section 3.9 alignment table to the rebuilt table and figure numbering."""
import docx
DOC='v5/thesis.docx'
d=docx.Document(DOC)
MAP={'Tables 1–5, Figures 5–10':'Tables 1–15, Figures 5–11',
     'Tables 6–18, Figures 11–24':'Tables 16–42, Figures 12–28',
     'Tables 19–24, Figures 25–30':'Tables 43–54, Figures 29–37'}
n=0
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for old,new in MAP.items():
                    if old in p.text and p.runs:
                        p.runs[0].text=p.text.replace(old,new)
                        for r in p.runs[1:]: r.text=''
                        n+=1
# the analysis column should now name the Kruskal-Wallis price work explicitly
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                if 'gross marketing margins and share of the end-of-chain price' in p.text and p.runs:
                    p.runs[0].text=p.text.replace(
                        'gross marketing margins and share of the end-of-chain price',
                        'gross marketing margins, the share of the end-of-chain price and the first-sale price spread by site')
                    for r in p.runs[1:]: r.text=''
                    n+=1
d.save(DOC); print('alignment-table cross-references updated:', n)
