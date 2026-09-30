# -*- coding: utf-8 -*-
"""Pull Chapters Four to Six out of the finished thesis, for the supervisors.

Taken from the thesis itself rather than rebuilt, so the formatting is the same
as the file the candidate is working in.
"""
import docx, re, shutil, sys
from docx.oxml.ns import qn

SRC = sys.argv[1] if len(sys.argv) > 1 else 'v8/thesis_trimmed.docx'
OUT = sys.argv[2] if len(sys.argv) > 2 else 'v8/Chapters_Four_to_Six.docx'
shutil.copy(SRC, OUT)

d = docx.Document(OUT)
body = d.element.body
def txt(el): return ''.join(t.text or '' for t in el.iter(qn('w:t'))).strip()
def heading(el):
    if el.tag != qn('w:p'): return ''
    pr = el.find(qn('w:pPr'))
    st = pr.find(qn('w:pStyle')) if pr is not None else None
    return (st.get(qn('w:val')) or '') if st is not None else ''

els = list(body)
start = next(i for i, el in enumerate(els)
             if heading(el).startswith('Heading') and txt(el).upper().startswith('CHAPTER FOUR'))
end = max(i for i, el in enumerate(els)
          if el.tag == qn('w:p') and txt(el).upper() == 'REFERENCES')

for el in els[:start] + els[end:]:
    if el.tag in (qn('w:p'), qn('w:tbl')):
        body.remove(el)

d.save(OUT)
k = docx.Document(OUT)
words = sum(len(p.text.split()) for p in k.paragraphs)
print(f'{OUT}: {len(k.paragraphs)} paragraphs, {len(k.tables)} tables, {words:,} paragraph words')
