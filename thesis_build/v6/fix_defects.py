# -*- coding: utf-8 -*-
"""Remove heading-styled paragraphs that carry no text.

The base document ends with an empty Heading 1. Word paints nothing for it, but
it takes a bookmark and a blank row in the table of contents, and it is the kind
of thing a reader notices before anything else.
"""
import docx
from docx.oxml.ns import qn

DOC = 'v6/thesis.docx'
d = docx.Document(DOC)
n = 0
for p in list(d.paragraphs):
    if (p.style.name or '').startswith('Heading') and not p.text.strip():
        p._p.getparent().remove(p._p)
        n += 1
# a section heading in the questionnaire had been pasted over itself
doubled = 0
for p in d.paragraphs:
    t = p.text.strip()
    if len(t) > 10 and len(t) % 2 == 0 and t[:len(t) // 2] == t[len(t) // 2:]:
        half = t[:len(t) // 2]
        if p.runs:
            p.runs[0].text = half
            for r in p.runs[1:]:
                r.text = ''
            doubled += 1

d.save(DOC)
print(f'empty heading paragraphs removed: {n}   doubled headings repaired: {doubled}')
