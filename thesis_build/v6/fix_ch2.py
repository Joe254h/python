# -*- coding: utf-8 -*-
"""Insert the critical assessment into Chapter Two and repair the two
sentences whose citations could not be verified."""
import docx, json, copy
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOC='v6/thesis.docx'
d=docx.Document(DOC)
ps=d.paragraphs

def retext(p, t):
    if p.runs:
        p.runs[0].text=t
        for r in p.runs[1:]: r.text=''

# ---- 1. the two unverifiable citations -----------------------------------
FIX=[('Mobile money can improve transaction security and create a basic record, '
      'although its use varies across coastal markets.',
      'Buyers may also advance cash or gear in return for continued or exclusive supply, '
      'an arrangement Crona et al. (2016) describe as solving a short-term financing problem '
      'while narrowing the seller’s choice of buyer. Whether mobile money changes that balance '
      'in this fishery is an empirical question taken up in Chapter Four.'),
     ('Kenya’s Blue Economy and sustainable-development commitments place greater emphasis on '
      'traceability, monitoring and value-chain development.',
      'Kenya’s Blue Economy and sustainable-development commitments place greater emphasis on '
      'traceability, monitoring and value-chain development (Fondo & Ogutu, 2021).')]
for old,new in FIX:
    for p in ps:
        if old in p.text:
            retext(p, p.text.replace(old,new)); print('fixed:', old[:58]); break
    else:
        print('NOT FOUND:', old[:58])

# ---- 2. rename 3.6 to the title the reviewer asked for --------------------
for p in ps:
    if p.text.strip().startswith('3.6 Sampling coverage and nonresponse') and p.runs:
        retext(p, '3.6 Non-Response and Sampling Limitations'); print('3.6 renamed'); break

# ---- 3. insert the critical assessment before the Theoretical Framework ---
BLK=json.load(open('v5/ch2.json'))
anchor=None
for p in d.paragraphs:
    if p.text.strip().endswith('Theoretical Framework') and (p.style.name or '').startswith('Heading'):
        anchor=p; break
assert anchor is not None, 'Theoretical Framework heading not found'

model_h2=None; model_h3=None; model_p=None
for p in d.paragraphs:
    st=p.style.name if p.style else ''
    if st=='Heading 2' and model_h2 is None: model_h2=p
    if st=='Heading 3' and model_h3 is None: model_h3=p
    if st=='Normal' and len(p.text.split())>25 and model_p is None: model_p=p
assert model_h2 is not None and model_p is not None

def clone(model, text):
    el=copy.deepcopy(model._p)
    first=None
    for t in el.iter(qn('w:t')):
        if first is None:
            first=t; t.text=text; t.set(qn('xml:space'),'preserve')
        else:
            t.text=''
    for n in el.findall(qn('w:pPr')) :
        for np in n.findall(qn('w:numPr')): n.remove(np)
    return el

prev=None
for b in BLK:
    model = model_h2 if b['k']=='h2' else (model_h3 if b['k']=='h3' else model_p)
    if model is None: model = model_p
    el=clone(model, b['t'])
    if prev is None: anchor._p.addprevious(el)
    else: prev.addnext(el)
    prev=el
print('Chapter Two critical assessment inserted:', len(BLK), 'blocks')
d.save(DOC)
