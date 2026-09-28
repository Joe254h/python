# -*- coding: utf-8 -*-
"""Rewrite the abstract so it matches the rebuilt results: four-actor tables,
the Kruskal-Wallis price comparisons and the value-distribution finding."""
import docx, copy
from docx.oxml.ns import qn
DOC = 'v6/thesis.docx'
d = docx.Document(DOC)
ps = d.paragraphs
i = next(k for k, p in enumerate(ps) if p.text.strip() == 'ABSTRACT')
target = None
for p in ps[i:i + 6]:
    if p.text.strip().startswith('Mud crab (Scylla serrata) fishing supports'):
        target = p; break
assert target is not None, 'abstract paragraph not found'

# drop any abstract paragraphs a previous run inserted, so this is idempotent
ti = next(k for k, p in enumerate(ps) if p._p is target._p)
while True:
    ps = d.paragraphs
    ti = next(k for k, p in enumerate(ps) if p._p is target._p)
    if ti + 1 >= len(ps):
        break
    nxt = ps[ti + 1].text.strip()
    if not nxt or nxt.startswith(('TABLE OF CONTENTS', 'DECLARATION', 'LIST OF')):
        break
    ps[ti + 1]._p.getparent().remove(ps[ti + 1]._p)

import sys
sys.path.insert(0, 'v6')
from abstract_text import PARAS

# first paragraph reuses the existing one; the rest are deep copies of it
target.runs[0].text = PARAS[0]
for r in target.runs[1:]:
    r.text = ''
prev = target._p
for text in PARAS[1:]:
    el = copy.deepcopy(target._p)
    first = None
    for t in el.iter(qn('w:t')):
        if first is None:
            first = t; t.text = text; t.set(qn('xml:space'), 'preserve')
        else:
            t.text = ''
    prev.addnext(el)
    prev = el
d.save(DOC)

d2 = docx.Document(DOC)
ps = d2.paragraphs
i = next(k for k, p in enumerate(ps) if p.text.strip() == 'ABSTRACT')
tot = 0
for p in ps[i + 1:i + 7]:
    t = p.text.strip()
    if not t or t == 'TABLE OF CONTENTS':
        continue
    print(f'[{len(t.split())} words] {t[:96]}...')
    tot += len(t.split())
print('abstract total:', tot, 'words')
