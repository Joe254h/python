# -*- coding: utf-8 -*-
"""Insert LIST OF ABBREVIATIONS after LIST OF FIGURES, and normalise KSh."""
import docx, copy
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

DOC = 'v6/thesis.docx'
d = docx.Document(DOC)
body = d.element.body
FONT = 'Times New Roman'

ROWS = [('BMU',      'Beach Management Unit'),
        ('df',       'Degrees of freedom'),
        ('FAO',      'Food and Agriculture Organization of the United Nations'),
        ('H',        'Kruskal–Wallis test statistic'),
        ('IBM SPSS', 'IBM Statistical Package for the Social Sciences'),
        ('IQR',      'Interquartile range (reported as the first to third quartile)'),
        ('IUCN',     'International Union for Conservation of Nature'),
        ('KeFS',     'Kenya Fisheries Service'),
        ('KMFRI',    'Kenya Marine and Fisheries Research Institute'),
        ('KSh',      'Kenya Shilling'),
        ('M',        'Mean'),
        ('Mdn',      'Median'),
        ('MSc',      'Master of Science'),
        ('N / n',    'Sample size / subgroup count'),
        ('p',        'Probability value (significance level)'),
        ('SCP',      'Structure–Conduct–Performance (framework)'),
        ('SD',       'Standard deviation'),
        ('WIO',      'Western Indian Ocean'),
        ('χ²',       'Chi-square test statistic')]

# ------------------------------------------------------------- idempotence
def ptx(el):
    return ''.join(n.text or '' for n in el.iter(qn('w:t'))).strip()

kids = list(body)
hdr = None
for i, el in enumerate(kids):
    if el.tag == qn('w:p') and ptx(el) == 'LIST OF ABBREVIATIONS':
        hdr = i
        break
if hdr is not None:                       # remove the old block wholesale
    j = hdr + 1
    while j < len(kids) and not (kids[j].tag == qn('w:p')
                                 and ptx(kids[j]).startswith('CHAPTER')):
        body.remove(kids[j]); j += 1
    body.remove(kids[hdr])
    print('previous abbreviations block removed')

# ------------------------------------------------------ find the insertion point
kids = list(body)
anchor = None
for i, el in enumerate(kids):
    if el.tag == qn('w:p') and ptx(el).startswith('CHAPTER ONE'):
        anchor = el
        break
assert anchor is not None, 'CHAPTER ONE not found'

def mkrun(text, *, bold=False):
    r = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
    rf = OxmlElement('w:rFonts')
    for a in ('w:ascii', 'w:hAnsi', 'w:cs'): rf.set(qn(a), FONT)
    rPr.append(rf)
    if bold: rPr.append(OxmlElement('w:b'))
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '24'); rPr.append(sz)
    col = OxmlElement('w:color'); col.set(qn('w:val'), '000000'); rPr.append(col)
    r.append(rPr)
    t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text
    r.append(t)
    return r

def mkpara(text, *, bold=False, before=0, after=60):
    p = OxmlElement('w:p'); pPr = OxmlElement('w:pPr')
    sp = OxmlElement('w:spacing')
    sp.set(qn('w:before'), str(before)); sp.set(qn('w:after'), str(after))
    sp.set(qn('w:line'), '276'); sp.set(qn('w:lineRule'), 'auto')
    pPr.append(sp); p.append(pPr)
    if text: p.append(mkrun(text, bold=bold))
    return p

# ---------------------------------------------------------------- heading
head = mkpara('LIST OF ABBREVIATIONS', bold=True, before=240, after=180)
anchor.addprevious(head)

# ---------------------------------------------------------------- the table
tbl = OxmlElement('w:tbl')
tblPr = OxmlElement('w:tblPr')
st = OxmlElement('w:tblStyle'); st.set(qn('w:val'), 'TableNormal'); tblPr.append(st)
w = OxmlElement('w:tblW'); w.set(qn('w:w'), '9026'); w.set(qn('w:type'), 'dxa'); tblPr.append(w)
lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
mar = OxmlElement('w:tblCellMar')
for side, v in (('top', '30'), ('left', '80'), ('bottom', '30'), ('right', '80')):
    e = OxmlElement('w:' + side); e.set(qn('w:w'), v); e.set(qn('w:type'), 'dxa'); mar.append(e)
tblPr.append(mar)
tbl.append(tblPr)
grid = OxmlElement('w:tblGrid')
for cw in (2100, 6926):
    gc = OxmlElement('w:gridCol'); gc.set(qn('w:w'), str(cw)); grid.append(gc)
tbl.append(grid)

def border(el, side, sz='8'):
    b = OxmlElement('w:' + side)
    b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), sz)
    b.set(qn('w:space'), '0'); b.set(qn('w:color'), '000000')
    el.append(b)

def mkrow(a, b, *, bold=False, top=False, bottom=False):
    tr = OxmlElement('w:tr')
    for txt, cw in ((a, 2100), (b, 6926)):
        tc = OxmlElement('w:tc'); tcPr = OxmlElement('w:tcPr')
        cwe = OxmlElement('w:tcW'); cwe.set(qn('w:w'), str(cw)); cwe.set(qn('w:type'), 'dxa')
        tcPr.append(cwe)
        bd = OxmlElement('w:tcBorders')
        if top: border(bd, 'top')
        if bottom: border(bd, 'bottom')
        if len(bd): tcPr.append(bd)
        cm = OxmlElement('w:tcMar')
        for side, v in (('top', '40'), ('left', '80'), ('bottom', '40'), ('right', '80')):
            e = OxmlElement('w:' + side); e.set(qn('w:w'), v); e.set(qn('w:type'), 'dxa'); cm.append(e)
        tcPr.append(cm)
        va = OxmlElement('w:vAlign'); va.set(qn('w:val'), 'center'); tcPr.append(va)
        tc.append(tcPr)
        p = OxmlElement('w:p'); pPr = OxmlElement('w:pPr')
        sp = OxmlElement('w:spacing')
        sp.set(qn('w:before'), '20'); sp.set(qn('w:after'), '20')
        sp.set(qn('w:line'), '240'); sp.set(qn('w:lineRule'), 'auto')
        pPr.append(sp); p.append(pPr)
        p.append(mkrun(txt, bold=bold))
        tc.append(p)
        tr.append(tc)
    return tr

tbl.append(mkrow('Abbreviation', 'Meaning', bold=True, top=True, bottom=True))
for k, (a, b) in enumerate(ROWS):
    tbl.append(mkrow(a, b, bottom=(k == len(ROWS) - 1)))
anchor.addprevious(tbl)
anchor.addprevious(mkpara('', after=120))

# ------------------------------------------------- normalise the KSh spelling
n = 0
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                if 'Ksh' in p.text and p.runs:
                    p.runs[0].text = p.text.replace('Ksh', 'KSh')
                    for r in p.runs[1:]: r.text = ''
                    n += 1
print('KSh spellings normalised:', n)

d.save(DOC)
print('LIST OF ABBREVIATIONS inserted:', len(ROWS), 'entries')
