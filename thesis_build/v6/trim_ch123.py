# -*- coding: utf-8 -*-
"""Tighten Chapters One to Three.

Chapters One to Three come from the base document rather than from a block
list, so they are edited on the assembled file. Paragraphs are matched on their
opening words. A replacement of None deletes the paragraph, which is how two
paragraphs are merged: the text of both goes into the first and the second is
removed. Every citation is kept.
"""
import docx, copy
from docx.oxml.ns import qn

DOC = 'v6/thesis.docx'
d = docx.Document(DOC)


def retext(p, t):
    """Put t in the paragraph, keeping the first run's character formatting."""
    runs = p._p.findall(qn('w:r'))
    if not runs:
        return False
    keep = runs[0]
    for r in runs[1:]:
        p._p.remove(r)
    for old in keep.findall(qn('w:t')):
        keep.remove(old)
    tn = copy.deepcopy(keep).makeelement(qn('w:t'), {})
    tn.set(qn('xml:space'), 'preserve')
    tn.text = t
    keep.append(tn)
    return True


EDITS = {
 # ---- 1.1 Background information: eight paragraphs to six
 'On the Kenyan coast, Scylla serrata is one of the commercially important':
  'On the Kenyan coast, Scylla serrata is one of the commercially important crustaceans '
  'available to small-scale fishers with relatively little equipment. Demand from hotels, '
  'local markets and export buyers has raised its commercial value, while reports of '
  'smaller crabs and the capture of immature animals suggest pressure on stocks and '
  'mangrove-dependent habitats. Kenyan research has examined the biology and ecology of mud '
  'crabs more closely than their markets, yet fishers also respond to the prices buyers '
  'offer, the grades buyers apply, access to alternative outlets, transport, credit and '
  'payment terms. When those conditions are poorly understood, management may address '
  'harvesting without addressing the incentives behind it.',
 'Kenyan research has examined the biology and ecology of mud crabs more closely': None,

 'The trade connects fishers, brokers, traders, retailers, hoteliers, exporters':
  'The trade connects fishers, brokers, traders, retailers, hoteliers, exporters and '
  'consumers through grading, pricing, storage, transport, packaging and credit. At many '
  'landing sites buyers grade and price crabs informally, and fishers who lack an '
  'independent price reference or another buyer have little room to challenge those '
  'decisions. Post-harvest handling matters because mud crabs are sold alive and their '
  'value depends on size, condition and survival: poor containment, unsuitable packaging '
  'and long journeys cause mortality or reduce quality, and those losses lower returns '
  'without reducing fishing effort.',
 'Post-harvest handling is especially important because mud crabs are usually sold': None,

 'Kenya’s Blue Economy agenda has renewed interest in mud crabs as a source':
  'Kenya’s Blue Economy agenda has renewed interest in mud crabs as a source of food, income '
  'and coastal employment. Crab fattening, community-based aquaculture and hatchery '
  'development are among the proposed alternatives to greater pressure on wild stocks, but '
  'their value depends on the market that already exists: who controls grades and prices, '
  'how crabs reach buyers, and where actors incur losses.',

 'Kwale County supports active mud crab harvesting in mangrove creeks':
  'Kwale County supports active mud crab harvesting in mangrove creeks and supplies traders, '
  'hotels and export-oriented markets. Even so, there is limited evidence on the roles of '
  'different actors, how grades and prices are set, how crabs are transported and packaged, '
  'the payment arrangements used, and the constraints at each market node. This study '
  'addresses that gap by examining the market structure and day-to-day business practices '
  'of the South Coast mud crab fishery.',

 # ---- 1.5 Significance: five paragraphs to four
 'For fishers and traders, the results provide a basis for clearer grades':
  'For fishers and traders, the results provide a basis for clearer grades, more transparent '
  'prices and better live-crab handling, and they show where transport, packaging and '
  'payment arrangements expose actors to avoidable loss or dependence. For BMUs, county '
  'officers and national fisheries agencies, the study offers actor-specific evidence for '
  'licensing support, training, infrastructure planning and communication about management '
  'measures; the cross-sectional data do not establish causal effects, but they show where '
  'practical interventions can be tested.',
 'For BMUs, county officers and national fisheries agencies, the study offers': None,

 # ---- 2.10.4: the one-line opener folded into the first limitation
 'Four limitations recur in the work reviewed above and shape what can be claimed': None,
 'The first is a production focus.':
  'Four limitations recur in the work reviewed above. The first is a production focus. Much '
  'of the Kenyan mud crab literature is biological, ecological or aquacultural: stock '
  'status, nursery habitat, growth and culture performance (Mirera, 2017a, 2017b; Mirera & '
  'Mosknes, 2015; Moser et al., 2005; Webley, 2008). That work is careful within its own '
  'terms but was not designed to describe exchange, so it cannot be read as evidence about '
  'prices, margins or bargaining.',

 # ---- 2.10.5
 'The literature specific to this coast is thinner than the regional literature':
  'The literature specific to this coast is thinner than the regional literature and is '
  'weighted towards production. Mirera et al. (2013) document fishing tactics and '
  'traditional knowledge among Kenyan mud crab fishers, including holding crabs at home for '
  'several days until enough have accumulated for sale, which makes survival part of the '
  'marketing problem rather than only a harvesting one. Ochiewo et al. (2010) and Mirera '
  '(2014a) record the gendered division of the coastal fishery, harvesting and first-tier '
  'trade male-dominated and women more visible in processing and trade. Fondo et al. (2020) '
  'describe the national status of the mangrove mud crab fishery and Fondo and Ogutu (2021) '
  'place it within the Blue Economy agenda, while Fulanda et al. (2009) and Kimani et al. '
  '(2018) supply the fisher population figures on which this study’s sampling frame rests.',

 'Kwale County itself appears in this literature mainly as a location':
  'Kwale County itself appears in this literature mainly as a location rather than as a '
  'subject. The one study treating its institutions directly is Njiru et al. (2021), who '
  'examine the socioeconomic factors influencing trader participation in Beach Management '
  'Units in Kwale; their finding that participation varies with the trader’s own '
  'circumstances supports communication designed for particular actor groups. Beyond that, '
  'what is known about Shimoni, Majoreni, Vanga and Msambweni as markets, rather than as '
  'landing sites, is largely undocumented.',

 'Three things follow. Prices at each node in a Kwale mud crab chain':
  'Three things follow. Prices at each node in a Kwale mud crab chain have not been '
  'reported; the grading criteria each actor category uses have not been recorded; and no '
  'study has tested whether actor characteristics, functions or constraints differ between '
  'the BMUs of this coast. Expectations formed at sector level cannot be assumed to hold at '
  'the level of a landing beach.',

 # ---- 2.10.6
 'The gap this study addresses is therefore specific rather than general.':
  'The gap this study addresses is specific rather than general. Previous work establishes '
  'that mud crab chains are functionally differentiated, that survival governs value and '
  'that small-scale harvesters are institutionally weak. It does not establish how the value '
  'of a South Coast mud crab is divided between the people who catch, trade and sell it, '
  'whether the grading rule behind that division is shared or asymmetric, or whether any of '
  'it varies between the Beach Management Units through which the fishery is co-managed. '
  'Those three questions are the three objectives of this study.',

 # ---- 2.12 Conceptual Framework: four paragraphs to three
 'Objective One describes market structure and actor profiles using actor category':
  'Objective One describes market structure and actor profiles using actor category, BMU, '
  'age, gender, education, experience, scale of operation, licensing, collective '
  'membership, credit and related business characteristics, summarised by actor category '
  'and compared across BMUs where coverage allowed. Objective Two examines conduct and '
  'functions at each node: harvesting and sourcing, grading, handling, packaging, '
  'transport, pricing, payment, buyer relations and market research. Reported income, '
  'prices, market access and post-harvest losses are presented as descriptive market '
  'outcomes rather than combined into a performance index.',
 'Objective Two examines conduct and functions at each node, including harvesting': None,

 'Objective Three describes the infrastructure, market barriers, policy awareness':
  'Objective Three describes the infrastructure, market barriers, policy awareness, risks '
  'and opportunities each actor category reported, read alongside the structure and conduct '
  'findings. No mediation or moderation model was estimated and the cross-sectional design '
  'does not support causal inference; actor category and BMU serve as grouping variables for '
  'description and for actor-specific association tests.',

 # ---- 3.2 Study site
 'The study was conducted on the South Coast of Kenya, mainly in Kwale County':
  'The study was conducted on the South Coast of Kenya, mainly in Kwale County, between '
  'latitudes 3.05 degrees and 4.75 degrees South and longitudes 38.52 degrees and 39.51 '
  'degrees East (Figure 4). The county covers about 8,270.2 square kilometres and includes a '
  'coastal strip from Likoni to Vanga, where mangrove forests, tidal creeks and shallow '
  'coastal waters provide breeding, feeding and shelter habitat for Scylla serrata and other '
  'species (Mirera, 2017a; Fondo et al., 2020). The South Coast was selected because it is '
  'an established harvesting and trading area. Small-scale fishers reach the mangroves on '
  'foot or by canoe using hand collection, baited traps and other simple gear, and '
  'year-round access makes mud crab fishing an important source of cash for some coastal '
  'households.',

 'The County has 20 Beach Management Units (BMUs) and 54 landing sites':
  'The county has 20 Beach Management Units (BMUs) and 54 landing sites, responsible for '
  'regulating fishing, managing landing sites and promoting sustainable use of fisheries '
  'resources. They are the entry points for understanding how mud crab harvesting and '
  'marketing work locally (Njiru et al., 2021). Landing sites are where fishers sell to '
  'brokers and traders, which makes them the right places to collect data on grading, '
  'pricing, transport, packaging and modes of payment.',

 'The South Coast is also strongly influenced by tourism and urban markets':
  'The South Coast is also strongly influenced by tourism and urban markets, which generate '
  'high demand for premium seafood such as live mud crabs. Hotels, restaurants and urban '
  'centres within and beyond the county provide ready markets, drawing multiple actors into '
  'the value chain (FAO, 2025), so mud crab fishing here has shifted from subsistence use '
  'towards a market-oriented activity.',

 # ---- 3.10.1
 'Questionnaires were administered on smartphones using KoBoToolbox':
  'Questionnaires were administered on smartphones using KoBoToolbox and KoBo Collect, and '
  'submitted records were reviewed, coded and downloaded in CSV and XLSX formats. The '
  'questionnaire followed the study objectives: the first section covered demographic and '
  'business characteristics, the second actor functions and market practices, the third '
  'constraints and opportunities.',
 'The questionnaire followed the study objectives. Its first section covered': None,

 'Descriptive analysis was conducted in IBM SPSS Statistics.':
  'Descriptive analysis was conducted in IBM SPSS Statistics. Categorical variables were '
  'summarised using frequencies and valid percentages within actor category. Continuous '
  'variables, including reported monthly income and mud crab prices, were summarised using '
  'the mean, standard deviation, median, first and third quartiles, minimum and maximum '
  'where valid observations were available. Every table in Chapter Four carries all four '
  'actor categories, so the four positions can be compared on the same row. Fishers and '
  'middlemen were additionally cross-tabulated by study site; hoteliers and exporters were '
  'not, because their samples were small and their site coverage uneven. All charts were '
  'generated from the same dataset.',

 # ---- 2.10.2 the middleman conflict
 'The clearest disagreement in this literature concerns whether middlemen':
  'The clearest disagreement in this literature is whether middlemen improve or reduce the '
  'position of fishers, and it cannot be settled by preference.',

 'One line of work treats the middleman as a net cost to the harvester.':
  'One line of work treats the middleman as a net cost to the harvester. Jacinto (2004) '
  'argues that in small-scale fisheries the share of final value captured at each node '
  'follows control of information and market access rather than physical effort, which '
  'implies that an actor positioned between the harvester and the buyer captures value the '
  'harvester cannot reach. Béné et al. (2007) extend this to credit: an advance against '
  'future supply solves an immediate cash problem while narrowing where the seller may sell, '
  'and they treat such arrangements as a mechanism of persistent low returns.',

 'A second line reaches almost the opposite conclusion.':
  'A second line reaches almost the opposite conclusion. Crona et al. (2016), working in '
  'Kenya and Zanzibar, describe middlemen as a critical social-ecological link whose '
  'relational knowledge is not easily replaced, and find them blamed for outcomes settled '
  'further along the chain while operating on modest turnovers and carrying the mortality '
  'risk. On that reading, removing the middleman would not transfer their margin to the '
  'fisher but would remove a function someone has to perform.',

 'The two positions are not reconcilable at the level of assertion':
  'The two positions are not reconcilable at the level of assertion, because they make '
  'different empirical claims: one about who captures the value, the other about what the '
  'middleman’s position returns once risk is accounted for. Resolving them needs the price '
  'at each node and the share each node holds, measured in the same chain. No study reports '
  'that for a Kenyan mud crab chain, which is the first gap this study addresses.',

 # ---- 2.10.3 grading
 'A second, less openly stated disagreement concerns grading.':
  'A second, less openly stated disagreement concerns grading. FAO (2025) treats grading by '
  'size, weight, shell condition and claw size as a legitimate quality signal supporting '
  'price formation in a live-crab trade. Jacinto (2004) treats the same practice as a site '
  'of informational asymmetry, because a criterion the buyer applies and the seller cannot '
  'verify transfers value rather than signalling it. Both may hold at once, but only if the '
  'rule is applied consistently and both parties can read it. Whether that condition holds '
  'has not been tested in a Kenyan mud crab market, and testing it means recording which '
  'criteria each actor category actually uses.',

 # ---- 2.11 Theoretical framework
 'This study uses the Structure-Conduct-Performance (SCP) paradigm':
  'This study uses the Structure-Conduct-Performance (SCP) paradigm as its main analytical '
  'framework. SCP is commonly applied in agricultural and fisheries marketing to organise '
  'evidence on market structure, actor conduct and market outcomes, and here it connects the '
  'composition and organisation of the mud crab market with the practices observed at each '
  'node. The framework guides interpretation; it is not used to claim a causal sequence.',

 'Market structure is described through the number and types of actors':
  'Market structure is described through the number and types of actors, their '
  'characteristics, entry conditions, licensing, access to capital, collective organisation '
  'and buyer connections. Conduct covers sourcing, grading, handling, pricing, payment, '
  'transport and credit relationships. Performance is considered through descriptive '
  'indicators available in the dataset: reported income and prices, market access, mortality '
  'or spoilage, and perceptions of market organisation. The constraints and opportunities '
  'reported under Objective Three give the operating context.',

 'Value-chain concepts are used as a mapping device':
  'Value-chain concepts are used as a mapping device to trace crabs from harvesting through '
  'aggregation, hospitality and export, not as a second causal theory. Livelihood research, '
  'in particular the sustainable livelihoods framework of Allison and Ellis (2001), provides '
  'background for interpreting education, experience, finance and collective membership, but '
  'the study measured no complete set of livelihood assets and estimated no livelihood '
  'effects. Keeping SCP central matches the descriptive design.',

 # ---- 3.5 sampling
 'The fisher sample was estimated from the available population information.':
  'The fisher sample was estimated from the available population information. Bowers (2017) '
  'reported mud crab fishers as about 3% of the 12,915 fishers along the Kenyan coast, and '
  'applying Kwale County’s estimated 26.9% share of the national mud crab fisher population '
  'gave a population of 104 fishers (Mirera, 2017a; Kimani et al., 2018).',

 'The completed survey covered 96 actors:':
  'The completed survey covered 96 actors: 65 mud crab fishers, 22 middlemen, five hoteliers '
  'and four exporters. Fishers were recruited from the selected BMUs using the available '
  'frame and field access. Because complete lists of downstream actors did not exist, '
  'respondents identified traders and buyers who could be approached, and those referrals '
  'were followed through snowball sampling.',

 # ---- 3.8 data collection
 'For Objective One, the questionnaire recorded actor category, age, sex':
  'For Objective One, the questionnaire recorded actor category, age, sex, marital status, '
  'education, household size, experience, monthly income, business ownership, scale of '
  'operation, collective membership, credit, licensing and institutional affiliation, which '
  'describe the people and businesses operating at each market node.',

 'For Objective Two, respondents described their harvesting or sourcing methods':
  'For Objective Two, respondents described their harvesting or sourcing methods, catch or '
  'traded volumes, grading, quality control, processing, preservation, packaging, transport, '
  'market destinations, price setting, payment, technology use, market research and sources '
  'of technical knowledge. Field observation checked reported handling, grading, transport '
  'and marketing practices where they could be seen.',

 'For Objective Three, respondents identified operational and market constraints':
  'For Objective Three, respondents identified operational and market constraints, '
  'infrastructure gaps, price risk, access to finance, technology, institutional support, '
  'regulation and capacity needs, and the opportunities they saw in market access, value '
  'addition, infrastructure, training, credit and institutional support. Key informants, '
  'including BMU leaders and government officers, gave context on sector-wide conditions; '
  'their accounts and the field notes were used to interpret the survey findings rather than '
  'as a separate causal test.',

 'Secondary evidence came from published studies, government reports':
  'Secondary evidence came from published studies, government reports and other documents on '
  'mud crab fisheries, actor roles and market conditions, and supplied the comparison '
  'literature used in Chapters One, Two and Five.',

 # ---- 3.10.4
 'Four measures summarise market structure.':
  'Four measures summarise market structure. Concentration at first sale is taken over the '
  'buying points fishers named, using the concentration ratio, the share of harvesters '
  'attached to the largest point, and the Herfindahl-Hirschman Index, the sum of the squared '
  'shares on a scale to 10,000; dividing 10,000 by that index gives the '
  'numbers-equivalent, the count of equal-sized outlets that would produce the same '
  'concentration. Because the survey recorded where fishers sold rather than what each buyer '
  'handled, these are shares of harvesters and not shares of volume. The gross marketing '
  'margin at a node is the difference between that actor’s selling and buying price as a '
  'percentage of his selling price; the total marketing margin applies the same formula '
  'across the chain; and the producer’s share is the fisher mean divided by the final-node '
  'mean. Physical mortality, the one cost the survey measured, is deducted from the volume '
  'the harvester realises, giving a producer’s share net of that loss. Price dispersion is '
  'the coefficient of variation, which lets nodes whose price levels differ by a factor of '
  'three be compared directly. All four were computed from the IBM SPSS Statistics output '
  'rather than by a built-in procedure, and the formulas are in the syntax file.',
}


def wc(p): return len(p.text.split())


before = sum(wc(p) for p in d.paragraphs)
done, deleted = set(), 0
for p in list(d.paragraphs):
    t = p.text.strip()
    if not t:
        continue
    for opener, new in EDITS.items():
        if opener in done or not t.startswith(opener):
            continue
        if new is None:
            p._p.getparent().remove(p._p)
            deleted += 1
        else:
            retext(p, new)
        done.add(opener)
        break

missing = [o[:56] for o in EDITS if o not in done]
if missing:
    print('NOT FOUND, nothing saved:')
    for m in missing:
        print('   ', m)
    raise SystemExit(1)

d.save(DOC)
after = sum(len(p.text.split()) for p in docx.Document(DOC).paragraphs)
print(f'Chapters One to Three: {len(done)} paragraphs edited, {deleted} merged away')
print(f'document paragraph words: {before:,} -> {after:,} (saved {before - after:,})')
